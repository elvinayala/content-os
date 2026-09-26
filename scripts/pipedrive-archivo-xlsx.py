"""Arma el Excel del archivo de Pipedrive a partir de los JSON que baja pipedrive-archivo.mjs.

uso: python3 scripts/pipedrive-archivo-xlsx.py <dir con los .json> <salida.xlsx>
Campos personalizados con su nombre real y opciones (enum/set) con su etiqueta; personas,
empresas, dueños, embudo y etapa legibles. Workbook en modo write_only (20K+ filas).
"""
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

d, salida = sys.argv[1], sys.argv[2]


def cargar(n):
    p = os.path.join(d, f"{n}.json")
    return json.load(open(p)) if os.path.exists(p) else []


usuarios = {u["id"]: u.get("name") for u in cargar("usuarios")}
embudos = {p["id"]: p.get("name") for p in cargar("embudos")}
etapas = {s["id"]: (s.get("name"), s.get("pipeline_id")) for s in cargar("etapas")}

ILEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def limpio(v):
    if v is None:
        return None
    if isinstance(v, (int, float, bool)):
        return v
    if isinstance(v, dict):
        # person_id/org_id/user_id vienen como {name, value, ...}
        if "name" in v:
            return limpio(v.get("name"))
        return limpio(json.dumps(v, ensure_ascii=False))
    if isinstance(v, list):
        partes = []
        for x in v:
            if isinstance(x, dict):
                partes.append(str(x.get("value") or x.get("name") or ""))
            else:
                partes.append(str(x))
        return limpio(", ".join(p for p in partes if p))
    s = ILEGAL.sub("", str(v))
    return s[:32000]  # tope de celda de Excel


def mapa_campos(campos):
    """key → (nombre, opciones{id: etiqueta})"""
    m = {}
    for c in campos:
        ops = {str(o["id"]): o.get("label") for o in (c.get("options") or [])}
        m[c["key"]] = (c.get("name") or c["key"], ops)
    return m


def valor_campo(v, ops):
    if v is None or v == "":
        return None
    if ops:
        ids = [x.strip() for x in str(v).split(",")] if not isinstance(v, list) else [str(x) for x in v]
        return ", ".join(ops.get(i, i) for i in ids)
    return limpio(v)


def hoja(wb, titulo, filas, campos=None, extras=None):
    ws = wb.create_sheet(titulo[:31])
    if not filas:
        ws.append(["(sin datos)"])
        return 0
    mc = mapa_campos(campos or [])
    claves = []
    for f in filas:
        for k in f.keys():
            if k not in claves:
                claves.append(k)
    cab = [mc[k][0] if k in mc else k for k in claves]
    ext = list((extras or {}).keys())
    enc = []
    for t in ext + cab:
        c = WriteOnlyCell(ws, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        enc.append(c)
    ws.append(enc)
    for f in filas:
        fila = [fn(f) for fn in (extras or {}).values()]
        for k in claves:
            v = f.get(k)
            fila.append(valor_campo(v, mc[k][1]) if k in mc else limpio(v))
        ws.append(fila)
    ws.freeze_panes = "A2"
    return len(filas)


wb = Workbook(write_only=True)
tratos = cargar("tratos")
n = {}
n["Tratos"] = hoja(
    wb,
    "Tratos",
    tratos,
    cargar("camposTrato"),
    {
        "Embudo": lambda f: embudos.get(f.get("pipeline_id")),
        "Etapa": lambda f: (etapas.get(f.get("stage_id")) or (None,))[0],
        "Dueño": lambda f: limpio(f.get("user_id")) or usuarios.get(f.get("user_id")),
    },
)
n["Personas"] = hoja(wb, "Personas", cargar("personas"), cargar("camposPersona"))
n["Empresas"] = hoja(wb, "Empresas", cargar("empresas"), cargar("camposEmpresa"))
n["Notas"] = hoja(wb, "Notas", cargar("notas"))
n["Actividades"] = hoja(wb, "Actividades", cargar("actividades"))
n["Leads (bandeja)"] = hoja(wb, "Leads (bandeja)", cargar("leads"))
n["Embudos y etapas"] = hoja(
    wb,
    "Embudos y etapas",
    cargar("etapas"),
    extras={"Embudo": lambda f: embudos.get(f.get("pipeline_id"))},
)
n["Usuarios"] = hoja(wb, "Usuarios", cargar("usuarios"))
campos = [dict(c, entidad=e) for e, cs in (("trato", cargar("camposTrato")), ("persona", cargar("camposPersona")), ("empresa", cargar("camposEmpresa"))) for c in cs]
n["Campos"] = hoja(wb, "Campos", campos)
wb.save(salida)
print("Hojas:", n)
